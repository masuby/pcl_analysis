"""What the call centre actually did with the leads we handed them.

Reads the two working sheets back — LBF AI Digital Agent Data and SME AI
Digital Agent Data — and reports how far each month's leads got: who was
called, who was reached, what they said, and who converted.

THE SHEETS DRIFT, AND THE REPORT HAS TO COPE
--------------------------------------------
These are live sheets people edit by hand, so what comes back is not always
what `distribute.py` wrote:

  * Agent names arrive in whatever case the agent typed. "ZULPHA" and "zulpha"
    are one person and are folded together — otherwise the same agent appears
    twice and their volume is halved.
  * Feedback is a dropdown, but free text still gets in ("not exist",
    "namba intumika kwasasa"). Anything off the list is kept, counted, and
    listed separately rather than silently dropped.
  * Columns move and disappear. The SME sheet currently has no Product or
    Assigned_to column at all, so lookups are by header name and a missing
    column reads as blank rather than throwing.

CONVERSIONS ARE REPORTED TWICE ON PURPOSE
-----------------------------------------
There are two independent signals — Feedback = "Converted" and the
`is_converted?` dropdown — and in the current data they disagree completely:
one row says Converted with is_converted = No, a different row says
is_converted = Yes with Feedback = Not interested. Reporting a single number
would be inventing certainty that is not there, so both are shown along with
how many rows disagree.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date

from .config import settings
from .tools.sheets import _services

PRODUCT_SHEETS = {
    "LBF": lambda: settings.lbf_sheet_id,
    "SME": lambda: settings.sme_sheet_id,
}

# The dropdown the call centre is meant to use.
FEEDBACK_OPTIONS = [
    "Not picking", "Not interested", "Request callback", "Not reachable",
    "Not qualified", "Request more time", "Failed to connect",
    "Duplicated number", "Qualified for SME", "Qualified for CS",
    "Converted", "Referred",
]
_CANON = {f.lower(): f for f in FEEDBACK_OPTIONS}

# Outcomes that mean a human actually spoke to the client. The rest are
# failures to make contact, which say nothing about whether the lead is any
# good — separating the two is the whole point of the report.
SPOKE_TO = {
    "not interested", "request callback", "not qualified", "request more time",
    "qualified for sme", "qualified for cs", "converted", "referred",
}
# Outcomes that mean the phone never connected.
NO_CONTACT = {
    "not picking", "not reachable", "failed to connect", "duplicated number",
}
# Genuine interest — worth chasing again.
WARM = {"request callback", "request more time", "qualified for sme",
        "qualified for cs", "converted", "referred"}


def _norm_feedback(v: str) -> tuple[str, bool]:
    """Return (canonical label, whether it was on the dropdown)."""
    s = " ".join(str(v or "").split())
    if not s:
        return "", True
    key = s.lower()
    if key in _CANON:
        return _CANON[key], True
    return s, False


def _norm_agent(v: str) -> str:
    """Fold an agent's name so case and spacing do not split one person in two."""
    s = " ".join(str(v or "").split())
    return s.title() if s else ""


def _pct(part: int, whole: int) -> float:
    return round(100.0 * part / whole, 1) if whole else 0.0


def _read_tab(svc, sheet_id: str, tab: str) -> tuple[list[str], list[list[str]]]:
    resp = svc.spreadsheets().values().get(
        spreadsheetId=sheet_id,
        range=f"'{tab.replace(chr(39), chr(39) * 2)}'!A1:Z100000",
    ).execute()
    values = resp.get("values", [])
    if not values:
        return [], []
    header = [str(h).strip() for h in values[0]]
    return header, values[1:]


def _column_reader(header: list[str]):
    """Look a column up by name, tolerating absence and short rows."""
    idx = {}
    for i, h in enumerate(header):
        key = h.strip().lower()
        if key and key not in idx:
            idx[key] = i

    def get(row: list[str], name: str) -> str:
        i = idx.get(name.strip().lower())
        if i is None or i >= len(row):
            return ""
        return str(row[i]).strip()

    return get, idx


def report_product(product: str, month: str = "") -> dict:
    """Build the report for one product's sheet."""
    sheet_id = PRODUCT_SHEETS[product]()
    if not sheet_id:
        return {"product": product, "ok": False,
                "error": f"AISM_{product}_SHEET_ID is not set, so its sheet cannot be read."}

    svc, _ = _services()
    meta = svc.spreadsheets().get(spreadsheetId=sheet_id).execute()
    title = meta["properties"]["title"]
    tabs = [s["properties"]["title"] for s in meta["sheets"]]

    wanted = [t for t in tabs if not month or month.strip().lower() in t.lower()]
    if not wanted:
        return {"product": product, "ok": False, "sheetTitle": title,
                "error": f"No tab matching {month!r}. Tabs present: {', '.join(tabs)}."}

    months = []
    for tab in wanted:
        header, rows = _read_tab(svc, sheet_id, tab)
        if not header:
            continue
        months.append(_summarise(product, title, tab, header, rows))

    return {"product": product, "ok": True, "sheetTitle": title,
            "sheetId": sheet_id, "tabs": tabs, "months": months}


def _summarise(product: str, title: str, tab: str,
               header: list[str], rows: list[list[str]]) -> dict:
    get, idx = _column_reader(header)
    has_agent = "assigned_to" in idx

    total = len(rows)
    worked = 0                    # a feedback was recorded
    off_list = Counter()
    feedback = Counter()
    by_agent = defaultdict(lambda: Counter())
    by_location = defaultdict(lambda: Counter())
    converted_feedback = []       # rows whose Feedback says Converted
    converted_flag = []           # rows whose is_converted? says Yes
    disagreements = []
    amounts = []
    comments = 0
    dates = Counter()

    for r in rows:
        raw_fb = get(r, "Feedback")
        label, known = _norm_feedback(raw_fb)
        agent = _norm_agent(get(r, "Assigned_to")) or "(unassigned)"
        location = get(r, "Location") or "(blank)"
        conv = get(r, "is_converted?").strip().lower()
        name = get(r, "Name")
        phone = get(r, "Phone Number")

        if get(r, "Date"):
            dates[get(r, "Date")] += 1
        if get(r, "Comments"):
            comments += 1

        amt = re.sub(r"[^\d.]", "", get(r, "Loan Amount"))
        if amt:
            try:
                amounts.append(float(amt))
            except ValueError:
                pass

        if label:
            worked += 1
            feedback[label] += 1
            if not known:
                off_list[label] += 1
            by_agent[agent][label] += 1
            by_location[location][label] += 1
        else:
            by_agent[agent]["(not worked)"] += 1
            by_location[location]["(not worked)"] += 1

        said_converted = label.lower() == "converted"
        flagged_yes = conv == "yes"
        if said_converted:
            converted_feedback.append({"name": name, "phone": phone, "agent": agent})
        if flagged_yes:
            converted_flag.append({"name": name, "phone": phone, "agent": agent,
                                   "feedback": label})
        if said_converted != flagged_yes and (said_converted or flagged_yes):
            disagreements.append({
                "name": name, "phone": phone, "agent": agent,
                "feedback": label or "(blank)",
                "isConverted": get(r, "is_converted?") or "(blank)",
            })

    spoke = sum(n for f, n in feedback.items() if f.lower() in SPOKE_TO)
    nocontact = sum(n for f, n in feedback.items() if f.lower() in NO_CONTACT)
    warm = sum(n for f, n in feedback.items() if f.lower() in WARM)

    agents = []
    for name, c in by_agent.items():
        w = sum(n for f, n in c.items() if f != "(not worked)")
        t = sum(c.values())
        s = sum(n for f, n in c.items() if f.lower() in SPOKE_TO)
        wm = sum(n for f, n in c.items() if f.lower() in WARM)
        agents.append({
            "agent": name, "assigned": t, "worked": w, "notWorked": t - w,
            "spokeTo": s, "warm": wm,
            "workedPct": _pct(w, t), "spokePct": _pct(s, w),
            "breakdown": dict(c.most_common()),
        })
    agents.sort(key=lambda a: (-a["assigned"], a["agent"]))

    locations = [
        {"location": loc,
         "leads": sum(c.values()),
         "worked": sum(n for f, n in c.items() if f != "(not worked)"),
         "spokeTo": sum(n for f, n in c.items() if f.lower() in SPOKE_TO),
         "warm": sum(n for f, n in c.items() if f.lower() in WARM)}
        for loc, c in by_location.items()
    ]
    locations.sort(key=lambda x: -x["leads"])

    notes = []
    if not has_agent:
        notes.append(
            "This sheet has no Assigned_to column, so the work cannot be broken "
            "down by agent — everything shows as unassigned.")
    if off_list:
        notes.append(
            "Feedback typed outside the dropdown: "
            + "; ".join(f"{k} ({v})" for k, v in off_list.most_common())
            + ". These are counted but do not roll up into any outcome.")
    if disagreements:
        notes.append(
            f"{len(disagreements)} row(s) where Feedback and is_converted? "
            "contradict each other, so the conversion count cannot be trusted "
            "until they are reconciled.")
    if worked and not amounts:
        notes.append("No loan amounts have been entered, so value cannot be reported.")
    if worked == 0:
        notes.append("Nobody has worked this list yet — no feedback has been recorded.")

    return {
        "tab": tab,
        "totals": {
            "leads": total,
            "worked": worked, "workedPct": _pct(worked, total),
            "notWorked": total - worked,
            "spokeTo": spoke, "spokePct": _pct(spoke, worked),
            "noContact": nocontact, "noContactPct": _pct(nocontact, worked),
            "warm": warm, "warmPct": _pct(warm, worked),
            "comments": comments,
            "loanAmountsEntered": len(amounts),
            "loanAmountTotal": round(sum(amounts), 2) if amounts else 0,
        },
        "conversions": {
            "byFeedback": len(converted_feedback),
            "byFlag": len(converted_flag),
            "agreeing": len([c for c in converted_feedback
                             if c["phone"] and any(f["phone"] == c["phone"]
                                                   for f in converted_flag)]),
            "feedbackRows": converted_feedback[:50],
            "flagRows": converted_flag[:50],
            "disagreements": disagreements[:50],
        },
        "feedback": [{"feedback": f, "count": n, "pct": _pct(n, worked),
                      "onDropdown": f.lower() in _CANON}
                     for f, n in feedback.most_common()],
        "agents": agents,
        "locations": locations[:25],
        "dates": dict(dates.most_common(10)),
        "notes": notes,
        "hasAgentColumn": has_agent,
    }


def build_report(month: str = "", products: list[str] | None = None) -> dict:
    """The whole report: every product, every matching month tab."""
    wanted = [p for p in (products or ["LBF", "SME"]) if p in PRODUCT_SHEETS]
    out = []
    for p in wanted:
        try:
            out.append(report_product(p, month))
        except Exception as exc:  # noqa: BLE001
            out.append({"product": p, "ok": False, "error": str(exc)})

    combined = {"leads": 0, "worked": 0, "spokeTo": 0, "warm": 0, "noContact": 0}
    for prod in out:
        for m in prod.get("months", []):
            t = m["totals"]
            for k in combined:
                combined[k] += t.get(k, 0)
    combined["workedPct"] = _pct(combined["worked"], combined["leads"])
    combined["spokePct"] = _pct(combined["spokeTo"], combined["worked"])
    combined["warmPct"] = _pct(combined["warm"], combined["worked"])

    return {"ok": True, "generatedAt": date.today().isoformat(),
            "month": month or "(all months)", "combined": combined,
            "products": out}


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

NAVY = "1F3864"


def _style_header(ws, ncols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _autofit(ws) -> None:
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(52, width + 3)


def build_workbook(month: str = "", products: list[str] | None = None) -> bytes:
    """The report as a workbook: one Summary sheet, then a sheet per product
    tab holding the outcome split, the agent table and the locations."""
    import io
    from openpyxl import Workbook

    rep = build_report(month, products)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Figure", "Value", "What it means"])
    c = rep["combined"]
    ws.append(["Leads distributed", c["leads"], "Rows handed to the call centre"])
    ws.append(["Worked", c["worked"], f"A feedback was recorded — {c['workedPct']}% of the list"])
    ws.append(["Not yet worked", c["leads"] - c["worked"], "No feedback recorded against them"])
    ws.append(["Spoke to the client", c["spokeTo"],
               f"{c['spokePct']}% of the worked rows actually reached a person"])
    ws.append(["Never connected", c["noContact"],
               "Not picking, not reachable, failed to connect or a duplicate number"])
    ws.append(["Genuinely interested", c["warm"],
               f"{c['warmPct']}% of worked rows — callback, more time, qualified, referred or converted"])
    ws.append([])
    ws.append(["Report generated", rep["generatedAt"], f"Month filter: {rep['month']}"])
    _style_header(ws, 3)
    _autofit(ws)

    for prod in rep["products"]:
        if not prod.get("ok"):
            s = wb.create_sheet(f"{prod['product']} (unavailable)"[:31])
            s.append(["Problem"])
            s.append([prod.get("error", "unknown")])
            _style_header(s, 1)
            _autofit(s)
            continue

        for m in prod["months"]:
            name = f"{prod['product']} {m['tab']}"[:31]
            s = wb.create_sheet(name)
            t = m["totals"]

            s.append(["Outcome", "Leads", "% of worked"])
            for f in m["feedback"]:
                label = f["feedback"] + ("" if f["onDropdown"] else "  (typed in, not on the dropdown)")
                s.append([label, f["count"], f["pct"]])
            if t["notWorked"]:
                s.append(["(not worked yet)", t["notWorked"], ""])
            s.append([])

            s.append(["Agent", "Assigned", "Worked", "Not worked",
                      "Spoke to", "Interested", "Worked %", "Spoke %"])
            agent_hdr = s.max_row
            for a in m["agents"]:
                s.append([a["agent"], a["assigned"], a["worked"], a["notWorked"],
                          a["spokeTo"], a["warm"], a["workedPct"], a["spokePct"]])
            s.append([])

            s.append(["Location", "Leads", "Worked", "Spoke to", "Interested"])
            loc_hdr = s.max_row
            for l in m["locations"]:
                s.append([l["location"], l["leads"], l["worked"], l["spokeTo"], l["warm"]])

            if m["notes"]:
                s.append([])
                s.append(["Worth knowing"])
                for n in m["notes"]:
                    s.append([n])

            _style_header(s, 3)
            from openpyxl.styles import Font
            for r in (agent_hdr, loc_hdr):
                for cell in s[r]:
                    if cell.value is not None:
                        cell.font = Font(bold=True)
            _autofit(s)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(month: str = "") -> str:
    tag = (month or date.today().strftime("%B")).strip().title()
    return f"AI_Leads_Callback_Report_{tag}_{date.today():%Y-%m-%d}.xlsx"
